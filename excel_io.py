"""Excel-lesing/-skriving via openpyxl.

Hele økten kjører på ett openpyxl Workbook som holdes åpent i minnet og lagres
til disk etter hver skanning. Vi bruker aldri data_only=True ved lasting -- det
ville permanent slettet eventuelle formler i fila ved neste lagring.

Rekkefølge ved enhver endring: muter alltid arket i minnet FØRST, forsøk deretter
å lagre. Slik går ingen endring tapt selv om selve disk-lagringen feiler (f.eks.
fordi fila er åpen i Excel) -- neste vellykkede lagring tar med seg alt som har
hopet seg opp i mellomtiden.
"""

import datetime
import os
import shutil

import openpyxl

MOTTAKSSTATUS_COLUMN = "Mottaksstatus"
UNMATCHED_COLUMN = "produsentnr ikke i produsentliste"


class LockedFileError(Exception):
    """Fila kunne ikke lagres fordi den er åpen/låst (f.eks. i Excel)."""


class UnsupportedFileType(Exception):
    pass


def normalize_cell(value) -> str:
    """openpyxl gir native typer (int/float/datetime), i motsetning til pandas'
    tvungne dtype=str. Normaliser alt til en sammenlignbar streng, og pass på at
    et heltall lagret som flyttall (11190453.0) ikke får en ".0"-hale."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ExcelSession:
    def __init__(self, path: str):
        ext = os.path.splitext(path)[1].lower()
        if ext not in (".xlsx", ".xlsm"):
            raise UnsupportedFileType(
                "Denne filen er i gammelt Excel-format (.xls) eller ukjent format.\n"
                "Lagre den som .xlsx i Excel og prøv igjen."
            )
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=False)
        self.ws = self.wb.active
        self.column_index: dict[str, int] = self._build_column_index()

        self.producer_col_idx: int | None = None
        self.analysis_col_idx: list[tuple[str, int]] = []
        self.sorting_col_idx: int | None = None
        self.mottaksstatus_col_idx: int | None = None
        self.unmatched_col_idx: int | None = None

        self._backup_done = False

    def _build_column_index(self) -> dict[str, int]:
        return {
            cell.value: cell.column
            for cell in self.ws[1]
            if cell.value not in (None, "")
        }

    def column_names(self) -> list[str]:
        return list(self.column_index.keys())

    # --- Kolonnevalg (fra wizard) ---

    def select_producer_column(self, name: str) -> None:
        self.producer_col_idx = self.column_index[name]

    def select_analysis_columns(self, names: list[str]) -> None:
        self.analysis_col_idx = [(name, self.column_index[name]) for name in names]

    def select_sorting_column(self, name: str | None) -> None:
        self.sorting_col_idx = self.column_index[name] if name else None

    # --- Sorteringsmodus: sikre at nødvendige kolonner finnes ---

    def _true_last_column(self) -> int:
        """Går bakover fra ws.max_column i header-raden til en ikke-tom celle,
        for å unngå overskyting pga. formaterte-men-tomme kolonner."""
        col = self.ws.max_column
        while col > 1 and (self.ws.cell(row=1, column=col).value in (None, "")):
            col -= 1
        return col

    def _ensure_column(self, header_name: str) -> int:
        if header_name in self.column_index:
            return self.column_index[header_name]
        new_col = self._true_last_column() + 1
        self.ws.cell(row=1, column=new_col, value=header_name)
        self.column_index[header_name] = new_col
        return new_col

    def ensure_sorting_columns(self) -> None:
        self.mottaksstatus_col_idx = self._ensure_column(MOTTAKSSTATUS_COLUMN)
        self.unmatched_col_idx = self._ensure_column(UNMATCHED_COLUMN)

    # --- Oppslag ---

    def _true_last_data_row(self, *col_indices: int) -> int:
        """Går bakover fra ws.max_row til en rad der minst én av de gitte
        kolonnene har innhold, for å finne riktig sted å søke/legge til en ny
        rad (unngår å lande midt i formaterte-men-tomme rader). Flere kolonner
        kan gis samtidig -- nødvendig fordi rader lagt til av
        register_unmatched_scan() har verdi i unmatched-kolonnen men ikke i
        produsentkolonnen, og motsatt for vanlige produsentrader."""
        row = self.ws.max_row
        while row > 1 and all(
            normalize_cell(self.ws.cell(row=row, column=c).value) == "" for c in col_indices
        ):
            row -= 1
        return row

    def find_matching_row(self, search_str: str) -> int | None:
        """Første rad (topp til bunn) der produsentkolonnen starter med search_str."""
        last_row = self._true_last_data_row(self.producer_col_idx)
        for row in range(2, last_row + 1):
            value = normalize_cell(self.ws.cell(row=row, column=self.producer_col_idx).value)
            if value.startswith(search_str):
                return row
        return None

    def get_producer_value(self, row: int) -> str:
        return normalize_cell(self.ws.cell(row=row, column=self.producer_col_idx).value)

    def get_analyses_to_perform(self, row: int) -> list[str]:
        result = []
        for name, col_idx in self.analysis_col_idx:
            if normalize_cell(self.ws.cell(row=row, column=col_idx).value) != "":
                result.append(name)
        return result

    def get_sorteringsparameter(self, row: int) -> int | None:
        """Returnerer 1-15, eller None hvis verdien mangler/er ugyldig."""
        if self.sorting_col_idx is None:
            return None
        raw = normalize_cell(self.ws.cell(row=row, column=self.sorting_col_idx).value)
        if raw.isdigit() and 1 <= int(raw) <= 15:
            return int(raw)
        return None

    # --- Sorteringsmodus: skriving ---

    def increment_mottaksstatus(self, row: int) -> int:
        raw = normalize_cell(self.ws.cell(row=row, column=self.mottaksstatus_col_idx).value)
        new_value = (int(raw) if raw.isdigit() else 0) + 1
        self.ws.cell(row=row, column=self.mottaksstatus_col_idx, value=new_value)
        return new_value

    def register_unmatched_scan(self, scanned_code: str) -> int:
        """Gjenbruker raden hvis koden allerede er registrert som ukjent tidligere
        i fila, og teller opp i Mottaksstatus-kolonnen (samme kolonne/betydning som
        for kjente produsenter: "hvor mange ganger mottatt"). Ellers legges en ny
        rad til nederst. Returnerer ny telling."""
        last_row = self._true_last_data_row(self.producer_col_idx, self.unmatched_col_idx)
        for row in range(2, last_row + 1):
            value = normalize_cell(self.ws.cell(row=row, column=self.unmatched_col_idx).value)
            if value == scanned_code:
                return self.increment_mottaksstatus(row)

        new_row = last_row + 1
        self.ws.cell(row=new_row, column=self.unmatched_col_idx, value=scanned_code)
        self.ws.cell(row=new_row, column=self.mottaksstatus_col_idx, value=1)
        return 1

    # --- Lagring ---

    def _make_backup(self) -> None:
        try:
            stem, ext = os.path.splitext(self.path)
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = f"{stem}.{ts}.bak{ext}"
            shutil.copy2(self.path, backup_path)
        except OSError:
            pass  # backup er et sikkerhetsnett, ikke en hard avhengighet

    def save(self) -> bool:
        """Lagrer arbeidsboka til disk. Returnerer False (uten å kaste) hvis fila
        er låst/åpen andre steder -- endringene forblir i minnet til neste
        vellykkede lagring."""
        if not self._backup_done:
            self._make_backup()
            self._backup_done = True
        try:
            self.wb.save(self.path)
            return True
        except PermissionError:
            return False
